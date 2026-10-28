"""
Analysis logic for the case conference prep tool.

Kept separate from the Streamlit UI so the UI can be swapped later
without touching the agent code.

Requires: pip install google-genai python-docx pandas openpyxl tenacity
Set your key: export GEMINI_API_KEY="AIzaSy..."
"""

import io
import os
from typing import Iterator, Optional

import pandas as pd
from google import genai
from google.genai import types
from google.genai.errors import APIError
from tenacity import retry, retry_if_exception_type, stop_after_attempt, wait_exponential

try:
    from docx import Document
except ImportError:
    Document = None

MODEL = "gemini-3.5-flash-lite"
MAX_TOKENS = 3000

SYSTEM_PROMPT = """You are an assistant helping Judge Xavier Rodriguez prepare for an early case management conference. Judge Rodriguez has a strong legal background and does not need exhaustive overviews of your reasoning; focus on accuracy and conciseness throughout your summary. Provide only the information that someone with a strong legal background would need to understand your recommendations and analysis. You read case documents and produce concise, neutral, well-organized summaries. 
 
Rules: 
- Use only what is in the documents. Do not invent facts, dates, parties, or citations. 
- Cite document names and specific page/paragraph numbers where possible. 
- If something is unclear, missing, or inconsistent, say so plainly. 
- Stay completely neutral between the parties."""

DOC_SUMMARY_PROMPT = """Summarize this document for a judge preparing for an early \
case management conference. Use this structure:

**Document type and purpose** (1-2 sentences)
**Key facts and claims** (bullets, with page/section references where possible)
**Relief or action requested** (if any)
**Deadlines, dates, or procedural posture mentioned**
**Gaps or ambiguities** (anything unclear or missing)

Be concise."""

OVERVIEW_PROMPT = """The documents for this case are provided below. Synthesize and condense all provided documents into one comprehensive pre-conference summary briefing for Judge Xavier Rodriguez.

Do NOT include recommendations, suggested questions, or judicial attention topics—provide ONLY the factual summary of the filings.

Use this structure:

**Summary**
A concise overview (3-5 sentences) summarizing the parties, nature of the suit, primary claims, defenses, and current procedural posture.

**Key Facts & Claims by Party**
Bullet points summarizing plaintiff allegations and defendant position/defenses with page/paragraph citations.

Write concisely and professionally for a federal district judge. Ensure every factual claim includes a precise document citation (e.g., Doc. 10, p. 2 ¶ 3)."""

TIMELINE_PROMPT = """Extract all significant dates, deadlines, and milestones mentioned across the provided documents and arrange them in strict chronological order.

For each entry, format the output in HTML matching this structure:

<div class="timeline-item">
    <div class="timeline-date">[Date, e.g., May 23, 2024 or By May 23, 2024]</div>
    <b>[Short Title / Description of Event or Deadline]</b><br>
    <span class="cite-tag">📄 [Exact Source, e.g., Doc. 10 - p. 2, ¶ 5]</span>
</div>

Rules:
- Order entries from earliest date to latest date.
- Include past historical events/filings AND proposed or future deadlines (e.g., discovery cutoffs, initial disclosures, motion deadlines, mediation, trial dates).
- Every item MUST include a precise citation to the source document, page, and paragraph/section where the date appears.
- Do NOT output any markdown headers, conversational prose, or standard lists—output ONLY the repeating `<div class="timeline-item">` blocks wrapped in a single parent container: `<div class="figma-card">...</div>`."""

DISCOVERY_PROMPT = """You are a judicial case management assistant helping a federal judge identify issues that may lead to early dismissal, judgment on the merits, settlement, or narrowing of a civil lawsuit. Your goal is to narrow the scope of proposed discovery leading from the provided documents as much as possible based on the parameters below, to assist in achieving speedier resolutions without missing any facts likely to be important for resolution accuracy. 

“Discovery” is the process of a lawyer, usually counsel for the plaintiff, making requests for documents or formal recorded interviews (“depositions”) from the other side, in hopes of clarifying factual uncertainty in the record. Review all documents to verify if discovery has been proposed or not. If discovery has been proposed, review the scope of proposed discovery (i.e., how many documents/depositions are being requested) and whether the initial amount of proposed discovery could be reduced. You should operate on a “tiered system” where the first tier should be the smallest possible scope of discovery that is most likely to answer unresolved open factual questions, such as (for example) only permitting depositions of the five names found in the documents most relevant to the dispute at issue or the selection of documents directly related to the case. Do not simply search for the most common names or documents in the case, instead analyze which pieces of missing information are most likely to answer questions that can dismiss or otherwise settle the case. 

The second tier should include auxiliary information that is possible but less likely to result in clarification of the factual disputes, such as depositions of all directly relevant parties or requests for a large swathe of documents that are likely to be relevant to the case. The third tier should be “broad net” approaches, such as requesting all documents that may be relevant or depositions of all involved parties.  

You should include all three tiers, but always try to suggest the smallest possible scope of discovery that has the greatest likelihood of answering factual issues key to the resolution of the case. Do not arbitrarily leave out depositions or discovery requests that are likely to produce valuable information, no matter how large the “First tier” is. Use your judgement and reference base from past case information to determine the smallest possible scope of “First tier” discovery (if any discovery is indeed necessary at all) that is likely to result in factual resolution of the case and suggest this three-tiered discovery plan based on your analysis.

Format the output clearly for Judge Xavier Rodriguez as follows:

**Current Discovery Posture**
Briefly state whether discovery has been proposed in the filings (e.g., in Rule 26(f) reports or initial disclosures), the requested scope/limits, and any reported disputes or agreements regarding ESI or depositions. Cite relevant document pages/paragraphs.

**Proposed Tiered Discovery Plan**

**Tier 1: Targeted Dispositive Discovery (Narrowest Scope / Highest Impact)**
Detail specific key depositions (named individuals) and narrowly targeted document requests most likely to resolve core factual questions that could lead to early dismissal, summary judgment, or settlement. Explain why each item is critical.

**Tier 2: Auxiliary Relevant Discovery (Moderate Scope)**
Detail secondary depositions and broader document categories that clarify secondary factual disputes if Tier 1 discovery does not fully resolve the lawsuit.

**Tier 3: Broad-Net Comprehensive Discovery (Full Scope)**
Outline standard or broad-net discovery requests (e.g., all relevant party depositions, extensive ESI searches) if full factual development across all claims and defenses becomes necessary."""

CONFERENCE_QUESTIONS_PROMPT = """You are a judicial case management assistant helping a federal judge identify issues that may lead to early dismissal, judgment on the merits, settlement, or narrowing of a civil lawsuit. You will receive case documents (e.g., complaint, answer, Rule 26(f) report) and an Excel spreadsheet dataset titled Common_Questions.xlsx provided below.

Step 1: Classify the Lawsuit. Review the uploaded documents and classify the case into the most appropriate categories: 

- Contract 
- IP - Patent Litigation 
- IP - Copyright 
- IP - Trade Secrets 
- IP - Trademarks 
- Civil Rights 
- Other Personal Injury 
- Real Property 
- Social Security 
- Motor Vehicle Personal Injury 
- Labor Suits 
- Tort Actions (non-personal injury) 
- Forfeitures and Penalties 

Select the primary category based on the claims asserted. If you identify multiple categories that the lawsuit falls under, or questions from other categories that may fit, you may consider those questions as well. 

Step 2: Retrieve Relevant Questions. Open the Common_Questions.xlsx dataset provided below and retrieve all questions under the column/section corresponding to the selected category. These questions identify issues that, if resolved, could potentially lead to early dismissal, judgment, settlement, or narrowing of the lawsuit. Evaluate every question in the selected column. Do not invent additional questions. 

Step 3: Evaluate Each Question. For each question, determine: 

1. Is the question relevant? Exclude questions that do not apply to the case. 

2. Do the documents address the question? Classify each as: 
- Uncontested: The documents answer the question without a material dispute. 
- Contested: The parties dispute the answer. 
- Unaddressed: The documents do not answer the question. 
- Partially Addressed: Some relevant information is available, but important details remain unresolved. 

3. Could resolving the question facilitate early resolution? 
- Uncontested issues: Flag if the undisputed facts could support early dismissal or judgment. 
- Contested issues: Generally do not flag if resolution requires substantial factual development. However, flag potentially dispositive legal issues that can be resolved without resolving disputed facts. 
- Unaddressed or partially addressed issues: Flag if obtaining the missing information could realistically facilitate early dismissal, judgment, settlement, or narrowing of claims. 

Do not flag issues merely because they are unanswered. Explain why their resolution could materially affect the litigation. 

Step 4: Produce the Report 

Case Classification: [Category and brief explanation] 

Potential Early Resolution Issues: 
For each flagged issue, provide: 
1. Question: The exact question from the spreadsheet. 
2. Status: Uncontested, Contested, Unaddressed, or Partially Addressed. 
3. Evidence: What the documents reveal, citing the document and page or paragraph. 
4. Significance: How resolving this issue could lead to early dismissal, judgment, settlement, or narrowing of the case. 
5. Recommended Action: A specific question the judge should ask or action to consider at the early case management conference. 

Rank flagged issues from most to least promising. 

Remaining Questions: Briefly list the questions not flagged and why they were excluded. 

General Rules: 
- Remain neutral and rely only on the uploaded documents for case-specific facts. 
- Cite the document and page, section, or paragraph supporting each factual conclusion. 
- Do not invent facts, citations, or answers to unresolved questions. 
- Distinguish allegations from established or uncontested facts. 
- Do not assume that an unanswered question establishes a legal defect. 
- Focus on practical opportunities for early resolution, not every possible legal issue. 
- If no questions present realistic opportunities for early resolution, say so. 

Goal: Identify the most promising issues from Common_Questions.xlsx that the judge could address early to facilitate dismissal, judgment, settlement, or narrowing of the lawsuit."""

DAMAGES_PROMPT = """Calculate potential damages from the uploaded lawsuit documents, using the provided Damages_Equations.xlsx dataset as a formula reference—not conclusive legal authority.  

Rules & Procedure:
1. Identify the relevant Case Type and matched formula(s) from Damages_Equations.xlsx corresponding to the claims asserted in the case documents.
2. Establish applicability: Identify claims, jurisdiction, relevant dates, and potentially available remedies. Flag uncertain legal prerequisites or ambiguous formulas for attorney review. Do not assume liability, treat demanded amounts as proven losses, or automatically apply maximum awards or multipliers.
3. Ground every input: Distinguish allegations, estimates, and established facts. Derive values only through transparent arithmetic from documented inputs; never invent values.
4. Handle missing information: Never substitute zero for an unknown. If an essential input is missing or unreadable, state “Not determinable from provided documents,” identify the missing input and evidence needed, and show the formula symbolically using plain standard text (DO NOT use LaTeX formulas like \\P_{lw}). Calculate supported components separately, labeling any subtotal incomplete, not total damages.
5. STRICT HTML TABLE FORMATTING REQUIREMENTS:
   - Output the table using CLEAN, VALID HTML (`<table>`, `<thead>`, `<tbody>`, `<tr>`, `<th>`, `<td>`). DO NOT use Markdown pipe `|` table syntax.
   - Wrap the entire table in `<div style="overflow-x: auto; width: 100%;">...</div>` for horizontal scrollability.
   - Use standard CSS inline styling for the table: `<table style="width:100%; border-collapse: collapse; min-width: 800px;">`.
   - Style headers with `background-color: #f8f9fa; text-align: left; padding: 8px 12px; border-bottom: 2px solid #dee2e6;`.
   - Style cells with `padding: 8px 12px; border-bottom: 1px solid #dee2e6; vertical-align: top;`.
   - Place extensive defense notes or legal citations in standard HTML bullet points (`<ul><li>...</li></ul>`) below the table.

Output Format:

<div style="overflow-x: auto; width: 100%;">
  <table style="width:100%; border-collapse: collapse; min-width: 800px;">
    <thead>
      <tr>
        <th style="background-color: #f8f9fa; text-align: left; padding: 10px; border-bottom: 2px solid #dee2e6;">Claim / Remedy</th>
        <th style="background-color: #f8f9fa; text-align: left; padding: 10px; border-bottom: 2px solid #dee2e6;">Formula</th>
        <th style="background-color: #f8f9fa; text-align: left; padding: 10px; border-bottom: 2px solid #dee2e6;">Inputs & Sources</th>
        <th style="background-color: #f8f9fa; text-align: left; padding: 10px; border-bottom: 2px solid #dee2e6;">Arithmetic / Result</th>
        <th style="background-color: #f8f9fa; text-align: left; padding: 10px; border-bottom: 2px solid #dee2e6;">Missing Info & Limitations</th>
      </tr>
    </thead>
    <tbody>
      <tr>
        <td style="padding: 10px; border-bottom: 1px solid #dee2e6; vertical-align: top;"><b>Claim Name</b></td>
        <td style="padding: 10px; border-bottom: 1px solid #dee2e6; vertical-align: top;">Formula text</td>
        <td style="padding: 10px; border-bottom: 1px solid #dee2e6; vertical-align: top;">Inputs & Sources</td>
        <td style="padding: 10px; border-bottom: 1px solid #dee2e6; vertical-align: top;">Result text</td>
        <td style="padding: 10px; border-bottom: 1px solid #dee2e6; vertical-align: top;">Limitations</td>
      </tr>
    </tbody>
  </table>
</div>

<br>
<h4>Notes, Defenses & Limitations</h4>
<ul>
  <li><b>Currency & Period:</b> Details here...</li>
  <li><b>Supported Subtotals:</b> Details here...</li>
  <li><b>Defenses & Positions:</b> Details here...</li>
</ul>"""


# ---------- Building model input from uploaded files ----------

def _docx_to_text(data: bytes) -> str:
    if Document is None:
        raise RuntimeError("python-docx is required to read Word files.")
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def build_content(name: str, data: bytes) -> list:
    """Convert an uploaded file into content parts accepted by Gemini."""
    ext = name.rsplit(".", 1)[-1].lower()

    if ext == "pdf":
        return [
            f'<document name="{name}">',
            types.Part.from_bytes(data=data, mime_type="application/pdf"),
            '</document>'
        ]

    if ext in ("png", "jpg", "jpeg"):
        mime_type = "image/png" if ext == "png" else "image/jpeg"
        return [
            f'<document name="{name}">',
            types.Part.from_bytes(data=data, mime_type=mime_type),
            '</document>'
        ]

    if ext == "docx":
        text = _docx_to_text(data)
    else:  # txt, md
        text = data.decode("utf-8", errors="replace")

    return [f'<document name="{name}">\n{text}\n</document>']


def load_common_questions(
    excel_bytes: Optional[bytes] = None, filepath: str = "Common_Questions.xlsx"
) -> str:
    """Reads Common_Questions.xlsx from bytes or disk and formats questions cleanly for prompt context."""
    try:
        if excel_bytes:
            df = pd.read_excel(io.BytesIO(excel_bytes), engine="openpyxl", header=None)
        elif os.path.exists(filepath):
            df = pd.read_excel(filepath, engine="openpyxl", header=None)
        else:
            return "<file name='Common_Questions.xlsx'>File not found on server or uploaded.</file>"

        case_type_row = None
        for r in range(len(df)):
            if df.iloc[r].astype(str).str.contains("Case Type").any():
                case_type_row = r
                break

        if case_type_row is None:
            df_std = pd.read_excel(
                io.BytesIO(excel_bytes) if excel_bytes else filepath,
                engine="openpyxl",
            )
            return f"<file name='Common_Questions.xlsx'>\n{df_std.to_string()}\n</file>"

        categories = df.iloc[case_type_row].values
        q_start_row = case_type_row + 3

        formatted_text = ["=== Common_Questions.xlsx Data ==="]

        for col_idx in range(1, len(categories)):
            cat_name = str(categories[col_idx]).strip()
            if not cat_name or cat_name == "nan":
                continue

            questions = df.iloc[q_start_row:, col_idx].dropna().tolist()
            questions = [
                str(q).strip()
                for q in questions
                if str(q).strip() and str(q).strip() != "nan"
            ]

            formatted_text.append(f"\n### Category: {cat_name}")
            for idx, q in enumerate(questions, 1):
                formatted_text.append(f"{idx}. {q}")

        return "\n".join(formatted_text)
    except Exception as e:
        return f"<file name='Common_Questions.xlsx'>Error parsing spreadsheet: {e}</file>"


def load_damages_equations(
    excel_bytes: Optional[bytes] = None, filepath: str = "Damages_Equations.xlsx"
) -> str:
    """Reads Damages_Equations.xlsx from bytes or disk and formats equations cleanly for prompt context."""
    try:
        if excel_bytes:
            df = pd.read_excel(io.BytesIO(excel_bytes), engine="openpyxl")
        elif os.path.exists(filepath):
            df = pd.read_excel(filepath, engine="openpyxl")
        else:
            return "<file name='Damages_Equations.xlsx'>Reference file not found on server or uploaded.</file>"

        formatted_text = ["=== Damages_Equations.xlsx Formulas Dataset ==="]
        
        for idx, row in df.iterrows():
            case_type = str(row.get('Case Type', '')).strip()
            damage_type = str(row.get('Type of Damages', '')).strip()
            formula = str(row.get('Formula', '')).strip()
            
            if case_type and case_type != "nan":
                formatted_text.append(
                    f"- Case Type: {case_type} | Category/Remedy: {damage_type} | Formula: {formula}"
                )

        return "\n".join(formatted_text)
    except Exception as e:
        return f"<file name='Damages_Equations.xlsx'>Error parsing damages spreadsheet: {e}</file>"


# ---------- Agent calls ----------

def get_client() -> genai.Client:
    """Reads GEMINI_API_KEY (or GOOGLE_API_KEY) from the environment."""
    return genai.Client()


@retry(
    stop=stop_after_attempt(5),
    wait=wait_exponential(multiplier=1, min=2, max=10),
    retry=retry_if_exception_type(APIError),
    reraise=True,
)
def _call_stream(client: genai.Client, contents: list):
    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        max_output_tokens=MAX_TOKENS,
    )
    return client.models.generate_content_stream(
        model=MODEL,
        contents=contents,
        config=config,
    )


def _stream(client: genai.Client, contents: list) -> Iterator[str]:
    response = _call_stream(client, contents)
    for chunk in response:
        if chunk.text:
            yield chunk.text


def stream_document_summary(client: genai.Client, name: str, data: bytes) -> Iterator[str]:
    """Stream a summary of a single document."""
    contents = build_content(name, data) + [DOC_SUMMARY_PROMPT]
    return _stream(client, contents)


def stream_case_overview(client: genai.Client, documents: list[tuple[str, bytes]]) -> Iterator[str]:
    """Passes all uploaded documents directly into Gemini to generate one consolidated summary briefing."""
    contents = []
    for name, data in documents:
        contents.extend(build_content(name, data))
    contents.append(OVERVIEW_PROMPT)
    return _stream(client, contents)


def stream_case_timeline(client: genai.Client, documents: list[tuple[str, bytes]]) -> Iterator[str]:
    """Extracts chronological dates and generates timeline HTML blocks."""
    contents = []
    for name, data in documents:
        contents.extend(build_content(name, data))
    contents.append(TIMELINE_PROMPT)
    return _stream(client, contents)


def stream_discovery_analysis(client: genai.Client, documents: list[tuple[str, bytes]]) -> Iterator[str]:
    """Generates a 3-tiered discovery scope evaluation based on case filings."""
    contents = []
    for name, data in documents:
        contents.extend(build_content(name, data))
    contents.append(DISCOVERY_PROMPT)
    return _stream(client, contents)


def stream_conference_questions(
    client: genai.Client,
    documents: list[tuple[str, bytes]],
    excel_bytes: Optional[bytes] = None,
) -> Iterator[str]:
    """Runs the conference questions evaluation using Common_Questions.xlsx and case documents."""
    contents = []
    for name, data in documents:
        contents.extend(build_content(name, data))

    excel_content = load_common_questions(excel_bytes=excel_bytes)
    contents.append(excel_content)
    contents.append(CONFERENCE_QUESTIONS_PROMPT)

    return _stream(client, contents)


def stream_damages_analysis(
    client: genai.Client,
    documents: list[tuple[str, bytes]],
    damages_excel_bytes: Optional[bytes] = None,
) -> Iterator[str]:
    """Calculates potential damages using case filings and Damages_Equations.xlsx formulas."""
    contents = []
    for name, data in documents:
        contents.extend(build_content(name, data))

    eq_content = load_damages_equations(excel_bytes=damages_excel_bytes)
    contents.append(eq_content)
    contents.append(DAMAGES_PROMPT)

    return _stream(client, contents)